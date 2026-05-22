"""
기상청 조건별통계 자동 수집 스크립트
- 지점: 제주(184), 서귀포(189), 성산(188), 고산(185)
- 요소: 기온, 강수량, 바람, 습도, 일조일사
- 기간: 당월 1일 ~ 어제
"""

import os
import io
import time
import requests
import pandas as pd
from datetime import date, timedelta
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# ── 지점 정보 ──────────────────────────────────────────────
STATIONS = {
    "제주":  "184",
    "서귀포": "189",
    "성산":  "188",
    "고산":  "185",
}

# ── 요소 코드 (기상청 파라미터) ────────────────────────────
ELEMENTS = {
    "기온":   "00",
    "강수량":  "01",
    "바람":   "02",
    "습도":   "03",
    "일조일사": "04",
}

# ── 출력 컬럼 순서 ─────────────────────────────────────────
COLS = [
    "일시",
    "평균기온(℃)", "최고기온(℃)", "최저기온(℃)",
    "강수량(mm)", "1시간최다강수량(mm)", "최다강수시각",
    "강우일수",
    "평균풍속(m/s)", "최대풍속(m/s)", "최대풍속시각",
    "최대순간풍속(m/s)", "최대순간풍속시각",
    "평균습도(%rh)", "최저습도(%rh)",
    "일조합(hr)", "일조율(%)", "일사합(MJ/m2)",
]

COL_WIDTHS = [12,11,11,11,11,16,14,10,12,12,14,14,16,12,12,11,10,13]

# ── 기상청 다운로드 URL ────────────────────────────────────
BASE_URL = "https://data.kma.go.kr/climate/RankState/selectRankStatisticsDivisionList.do"
DOWNLOAD_URL = "https://data.kma.go.kr/climate/RankState/downloadRankStatisticsDivisionList.do"


def get_period():
    """당월 1일 ~ 어제"""
    today = date.today()
    start = today.replace(day=1)
    end   = today - timedelta(days=1)
    return start.strftime("%Y%m%d"), end.strftime("%Y%m%d")


def fetch_xls(stn_code: str, elem_code: str, start: str, end: str) -> bytes | None:
    """기상청에서 XLS 파일 다운로드 (cp949 탭 구분 텍스트 형식)"""
    session = requests.Session()
    session.headers.update({
        "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36",
        "Referer": BASE_URL,
    })

    params = {
        "pgmNo": "179",
        "stnIds": stn_code,
        "startDt": start,
        "endDt": end,
        "elementCd": elem_code,
        "divisionCd": "S",  # 일별
    }

    try:
        # 세션 초기화
        session.get(BASE_URL, params={"pgmNo": "179"}, timeout=15)
        time.sleep(0.5)

        resp = session.get(DOWNLOAD_URL, params=params, timeout=30)
        resp.raise_for_status()

        if len(resp.content) < 100:
            return None
        return resp.content
    except Exception as e:
        print(f"  ⚠ 다운로드 실패 stn={stn_code} elem={elem_code}: {e}")
        return None


def parse_xls(raw: bytes) -> pd.DataFrame:
    """cp949 탭 구분 텍스트 → DataFrame"""
    content = raw.decode("cp949", errors="replace")
    lines = content.split("\r\n")

    # 헤더 행 찾기
    header_idx = None
    for i, line in enumerate(lines):
        if "지점번호" in line or "일시" in line:
            header_idx = i
            break

    if header_idx is None:
        return pd.DataFrame()

    df = pd.read_csv(
        io.StringIO("\r\n".join(lines[header_idx:])),
        sep="\t",
        dtype=str,
    )
    df = df.dropna(subset=["일시"])
    df = df[df["일시"].str.match(r"\d{4}-\d{2}-\d{2}", na=False)]
    df = df.reset_index(drop=True)
    return df


def build_station_df(stn_name: str, stn_code: str, start: str, end: str) -> pd.DataFrame:
    """한 지점의 모든 요소를 합쳐 최종 DataFrame 반환"""
    dfs = {}
    for elem_name, elem_code in ELEMENTS.items():
        print(f"  📥 {stn_name} - {elem_name} 수집 중...")
        raw = fetch_xls(stn_code, elem_code, start, end)
        if raw:
            df = parse_xls(raw)
            if not df.empty:
                dfs[elem_name] = df
                print(f"     ✅ {len(df)}행")
            else:
                print(f"     ⚠ 파싱 결과 없음")
        time.sleep(0.3)

    if not dfs:
        return pd.DataFrame(columns=COLS)

    # 날짜 기준 DataFrame 선택 (기온 우선, 없으면 첫 번째)
    base_key = "기온" if "기온" in dfs else list(dfs.keys())[0]
    base = dfs[base_key][["일시"]].copy()

    # ── 기온 ──
    if "기온" in dfs:
        t = dfs["기온"]
        for col in ["평균기온(℃)", "최고기온(℃)", "최저기온(℃)"]:
            c = next((c for c in t.columns if col.replace("(", "").replace(")", "").replace("℃","") in c), None)
            if c:
                base = base.merge(t[["일시", c]].rename(columns={c: col}), on="일시", how="left")

    # ── 강수량 ──
    if "강수량" in dfs:
        r = dfs["강수량"]
        col_map = {}
        for orig, new in [("강수량(mm)", "강수량(mm)"),
                          ("1시간최다강수량(mm)", "1시간최다강수량(mm)"),
                          ("1시간최다강수량시각", "최다강수시각")]:
            c = next((c for c in r.columns if orig.split("(")[0] in c), None)
            if c:
                col_map[c] = new
        if col_map:
            base = base.merge(
                r[["일시"] + list(col_map.keys())].rename(columns=col_map),
                on="일시", how="outer"
            ).sort_values("일시").reset_index(drop=True)

    # ── 바람 ──
    if "바람" in dfs:
        w = dfs["바람"]
        col_map = {}
        for orig, new in [("평균풍속(m/s)", "평균풍속(m/s)"),
                          ("최대풍속(m/s)", "최대풍속(m/s)"),
                          ("최대풍속시각", "최대풍속시각"),
                          ("최대순간풍속(m/s)", "최대순간풍속(m/s)"),
                          ("최대순간풍속시각", "최대순간풍속시각")]:
            c = next((c for c in w.columns if orig.split("(")[0].replace("시각","") in c), None)
            if c:
                col_map[c] = new
        if col_map:
            base = base.merge(
                w[["일시"] + list(col_map.keys())].rename(columns=col_map),
                on="일시", how="left"
            )

    # ── 습도 ──
    if "습도" in dfs:
        h = dfs["습도"]
        col_map = {}
        for orig, new in [("평균습도(%rh)", "평균습도(%rh)"),
                          ("최저습도(%rh)", "최저습도(%rh)")]:
            c = next((c for c in h.columns if orig.split("(")[0] in c), None)
            if c:
                col_map[c] = new
        if col_map:
            base = base.merge(
                h[["일시"] + list(col_map.keys())].rename(columns=col_map),
                on="일시", how="left"
            )

    # ── 일조일사 ──
    if "일조일사" in dfs:
        s = dfs["일조일사"]
        col_map = {}
        for orig, new in [("일조합(hr)", "일조합(hr)"),
                          ("일조율(%)", "일조율(%)"),
                          ("일사합(MJ/m2)", "일사합(MJ/m2)")]:
            c = next((c for c in s.columns if orig.split("(")[0] in c), None)
            if c:
                col_map[c] = new
        if col_map:
            base = base.merge(
                s[["일시"] + list(col_map.keys())].rename(columns=col_map),
                on="일시", how="left"
            )

    # ── 강우일수 계산 ──
    if "강수량(mm)" in base.columns:
        base["강우일수"] = (pd.to_numeric(base["강수량(mm)"], errors="coerce") > 0).astype(int)
    else:
        base["강우일수"] = ""

    # 없는 컬럼 빈칸으로 채우기
    for col in COLS:
        if col not in base.columns:
            base[col] = ""

    return base[COLS]


def save_excel(station_data: dict[str, pd.DataFrame], output_path: str):
    """지점별 시트로 구성된 엑셀 저장"""
    wb = Workbook()
    wb.remove(wb.active)

    header_fill = PatternFill("solid", fgColor="4472C4")
    header_font = Font(color="FFFFFF", bold=True, size=10)
    title_font  = Font(bold=True, size=12)
    thin = Side(style="thin", color="CCCCCC")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for stn_name, df in station_data.items():
        ws = wb.create_sheet(title=stn_name)

        # 1행: 지점명 타이틀
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(COLS))
        title_cell = ws.cell(row=1, column=1, value=f"[{stn_name}] 기상 관측 데이터")
        title_cell.font = title_font
        title_cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[1].height = 22

        # 2행: 헤더
        for c_idx, col_name in enumerate(COLS, 1):
            cell = ws.cell(row=2, column=c_idx, value=col_name)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = border
        ws.row_dimensions[2].height = 30

        # 3행~: 데이터
        for r_idx, row in df.iterrows():
            for c_idx, col_name in enumerate(COLS, 1):
                val = row[col_name]
                if val == "" or pd.isna(val):
                    val = None
                cell = ws.cell(row=r_idx + 3, column=c_idx, value=val)
                cell.alignment = Alignment(horizontal="center", vertical="center")
                cell.border = border

        # 컬럼 너비
        for c_idx, width in enumerate(COL_WIDTHS, 1):
            ws.column_dimensions[get_column_letter(c_idx)].width = width

        # 틀 고정 (A3)
        ws.freeze_panes = "A3"

    wb.save(output_path)
    print(f"\n✅ 저장 완료: {output_path}")


def main():
    start, end = get_period()
    today_str = date.today().strftime("%Y%m%d")
    month_str = date.today().strftime("%Y%m")

    print(f"🌤 제주 기상 데이터 수집 시작")
    print(f"   기간: {start} ~ {end}")
    print(f"   지점: {', '.join(STATIONS.keys())}\n")

    os.makedirs("data", exist_ok=True)
    output_path = f"data/jeju_weather_{month_str}.xlsx"

    station_data = {}
    for stn_name, stn_code in STATIONS.items():
        print(f"\n📍 {stn_name} ({stn_code})")
        df = build_station_df(stn_name, stn_code, start, end)
        station_data[stn_name] = df
        print(f"   → {len(df)}행 완성")

    save_excel(station_data, output_path)

    # GitHub Actions 출력 변수
    with open(os.environ.get("GITHUB_OUTPUT", "/dev/null"), "a") as f:
        f.write(f"output_file={output_path}\n")
        f.write(f"month={month_str}\n")


if __name__ == "__main__":
    main()
